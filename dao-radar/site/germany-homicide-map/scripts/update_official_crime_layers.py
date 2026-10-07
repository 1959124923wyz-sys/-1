#!/usr/bin/env python3
from __future__ import annotations
import hashlib, io, json, os, re, sqlite3, tempfile, time
from datetime import datetime, timezone
from pathlib import Path

import requests
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
OUT_NATIONAL=ROOT/"data/germany_official_violence.json"
OUT_BERLIN=ROOT/"data/berlin_official_violence.geojson"

PKS_DB="https://stadtritter.de/wp-content/plugins/stadtritter-krimstatistik/data/pks-historie.db"
PKS_MANIFEST="https://stadtritter.de/wp-content/plugins/stadtritter-krimstatistik/data/pks-historie-manifest.json"
BERLIN_XLSX="https://www.berlin.de/polizei/_assets/dienststellen/lka/fallzahlen_hz-2016-2025.xlsx?ts=1790656142"
BERLIN_WFS="https://gdi.berlin.de/services/wfs/lor_2021"
BERLIN_WFS_PARAMS={
    "service":"wfs","version":"2.0.0","request":"GetFeature",
    "typeNames":"lor_2021:b_lor_bzr_2021",
    "outputFormat":"application/json","srsName":"EPSG:4326"
}
HEAD={"User-Agent":"GermanyCrimeMonitorOfficialLayers/1.0 (+https://github.com/1959124923wyz-sys/-1)"}

def save(path,obj):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")

def get(session,url,*,params=None,timeout=90):
    last=None
    for attempt in range(5):
        r=session.get(url,params=params,headers=HEAD,timeout=timeout)
        last=r
        if r.status_code==429:
            raw=r.headers.get("Retry-After","")
            try: wait=float(raw)
            except Exception: wait=min(45,3*(2**attempt))
            time.sleep(max(2,min(wait,60))); continue
        if 500<=r.status_code<600:
            time.sleep(min(30,2*(2**attempt))); continue
        r.raise_for_status(); return r
    if last is not None:last.raise_for_status()
    raise RuntimeError(url)

def norm_key(v,width=None):
    s=re.sub(r"\D","",str(v or ""))
    if not s:return ""
    if width:s=s.zfill(width)
    return s

def prop(props,*names):
    nm={re.sub(r"[^A-Z0-9]","",str(k).upper()):v for k,v in props.items()}
    for name in names:
        key=re.sub(r"[^A-Z0-9]","",name.upper())
        if key in nm and nm[key] not in (None,""):return nm[key]
    return None

session=requests.Session()

# --- Germany: BKA PKS 2025 Gewaltkriminalität by state -----------------------
manifest=get(session,PKS_MANIFEST).json()
db_bytes=get(session,PKS_DB,timeout=120).content
want_sha=(manifest.get("database") or {}).get("sha256")
got_sha=hashlib.sha256(db_bytes).hexdigest()
if want_sha and want_sha!=got_sha:
    raise RuntimeError(f"PKS mirror sha mismatch: {got_sha} != {want_sha}")

fd,db_path=tempfile.mkstemp(suffix=".db"); os.close(fd)
Path(db_path).write_bytes(db_bytes)
con=sqlite3.connect(db_path)
try:
    source=con.execute("""
      SELECT anbieter,dateiname,tabellennummer,version,download_revision,url,sha256,population_basis,coverage_scope
      FROM quellen WHERE jahr=2025 AND ebene='land'
    """).fetchone()
    if not source or source[0]!="Bundeskriminalamt":
        raise RuntimeError(f"Unexpected PKS source: {source}")
    rows=con.execute("""
      SELECT r.region_id,r.name,f.faelle,f.hz,d.delikt_name
      FROM fakten f
      JOIN regionen r ON r.jahr=f.jahr AND r.ebene=f.ebene AND r.region_id=f.region_id
      JOIN delikte d ON d.jahr=f.jahr AND d.ebene=f.ebene AND d.delikt_schluessel=f.delikt_schluessel
      WHERE f.jahr=2025 AND f.ebene='land' AND f.delikt_schluessel='892000'
      ORDER BY r.region_id
    """).fetchall()
finally:
    con.close()
    try:os.remove(db_path)
    except OSError:pass

if len(rows)!=16:
    raise RuntimeError(f"Expected 16 state Gewaltkriminalität rows, got {len(rows)}")

states={}
for region_id,name,cases,hz,delikt_name in rows:
    if cases is None or hz is None:continue
    states[name]={
        "ags":str(region_id),"cases":int(cases),"hz":round(float(hz),2),
        "offense_key":"892000","offense":str(delikt_name)
    }
if len(states)!=16:raise RuntimeError(f"Only {len(states)} usable state rows")

national={
  "meta":{
    "generated_at":datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00","Z"),
    "year":2025,
    "metric":"Gewaltkriminalität",
    "metric_key":"892000",
    "unit":"cases per 100,000 inhabitants (HZ)",
    "scope":"16 German states",
    "primary_source":"Bundeskriminalamt, Polizeiliche Kriminalstatistik 2025",
    "primary_source_url":source[5],
    "transport_mirror":"Stadtritter PKS open-data mirror",
    "transport_url":PKS_DB,
    "transport_manifest":PKS_MANIFEST,
    "transport_sha256":got_sha,
    "source_file_sha256":source[6],
    "source_table":source[2],
    "source_version":source[3],
    "source_revision":source[4],
    "population_basis":source[7],
    "note":"Annual structural background. HZ values are the official PKS population-standardized frequency values and are not a 90-day live count."
  },
  "states":states
}
save(OUT_NATIONAL,national)
print("NATIONAL",len(states),"min",min(x["hz"] for x in states.values()),"max",max(x["hz"] for x in states.values()))

# --- Berlin: 2025 Kriminalitätsatlas BZR official violence components -------
xlsx=get(session,BERLIN_XLSX,timeout=90).content
wb=openpyxl.load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
if "HZ_2025" not in wb.sheetnames or "Fallzahlen_2025" not in wb.sheetnames:
    raise RuntimeError(wb.sheetnames)

def read_sheet(sheet):
    ws=wb[sheet]
    headers=[str(x or "").replace("\n"," ").strip() for x in next(ws.iter_rows(min_row=5,max_row=5,values_only=True))]
    out={}
    for row in ws.iter_rows(min_row=6,values_only=True):
        raw=row[0]
        code=norm_key(raw,6)
        name=str(row[1] or "").strip()
        # Bezirksregion rows are six-digit LOR keys, exclude district totals and unassigned buckets.
        if len(code)!=6 or code.endswith("0000") or code.endswith("9900") or not name:
            continue
        vals={headers[i]:row[i] for i in range(min(len(headers),len(row)))}
        out[code]={"name":name,"values":vals}
    return headers,out

hz_headers,hz_rows=read_sheet("HZ_2025")
case_headers,case_rows=read_sheet("Fallzahlen_2025")

def find_col(headers,pattern):
    rx=re.compile(pattern,re.I)
    for h in headers:
        if rx.search(h):return h
    raise RuntimeError(f"column not found: {pattern}; headers={headers}")

hz_rob=find_col(hz_headers,r"^Raub$")
hz_severe=find_col(hz_headers,r"Gefährl.*schwere.*Körper")
case_rob=find_col(case_headers,r"^Raub$")
case_severe=find_col(case_headers,r"Gefährl.*schwere.*Körper")

geo=get(session,BERLIN_WFS,params=BERLIN_WFS_PARAMS,timeout=120).json()
features=geo.get("features") or []
if len(features)<100:
    raise RuntimeError(f"Expected Berlin BZR features, got {len(features)}")

out_features=[];matched=0
for f in features:
    props=f.get("properties") or {}
    raw=prop(props,"BZR_ID","BZR","RAUMID","SCHLUESSEL","SCHLÜSSEL")
    code=norm_key(raw,6)
    if code not in hz_rows:
        # Some services expose an 8-digit PLR-like key; try the first 6 digits.
        digits=norm_key(raw)
        code=digits[:6] if len(digits)>=6 else code
    h=hz_rows.get(code); c=case_rows.get(code)
    if not h or not c:continue
    try:
        rob_hz=float(h["values"][hz_rob]); sev_hz=float(h["values"][hz_severe])
        rob_cases=int(c["values"][case_rob]); sev_cases=int(c["values"][case_severe])
    except Exception:
        continue
    composite=rob_hz+sev_hz
    p=dict(props)
    p.update({
      "bzg_id":code,
      "bzg_name":h["name"],
      "year":2025,
      "robbery_hz":round(rob_hz,2),
      "severe_injury_hz":round(sev_hz,2),
      "violence_proxy_hz":round(composite,2),
      "robbery_cases":rob_cases,
      "severe_injury_cases":sev_cases,
      "violence_proxy_cases":rob_cases+sev_cases,
      "metric_note":"Derived sum of official 2025 PKS Raub + Gefährliche/schwere Körperverletzung HZ; not identical to BKA Gewaltkriminalität."
    })
    out_features.append({"type":"Feature","geometry":f.get("geometry"),"properties":p})
    matched+=1

if matched<120:
    sample_props=features[0].get("properties") if features else {}
    raise RuntimeError(f"Only {matched} Berlin BZR matches; sample props={sample_props}")

berlin={
  "type":"FeatureCollection",
  "meta":{
    "generated_at":datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00","Z"),
    "year":2025,
    "scope":"Berlin Bezirksregionen",
    "feature_count":matched,
    "primary_source":"Polizei Berlin, Kriminalitätsatlas 2025",
    "source_url":"https://www.berlin.de/polizei/service/kriminalitaetsatlas/",
    "source_xlsx":BERLIN_XLSX,
    "geometry_source":"Geoportal Berlin LOR 2021 WFS",
    "geometry_url":BERLIN_WFS,
    "default_metric":"violence_proxy_hz",
    "metrics":{
       "violence_proxy_hz":"Raub + gefährliche/schwere Körperverletzung, HZ per 100,000 (derived sum)",
       "severe_injury_hz":"Gefährliche und schwere Körperverletzung, HZ per 100,000",
       "robbery_hz":"Raub, HZ per 100,000"
    },
    "note":"Annual structural background; the combined metric is a transparent derived sum of two non-identical official PKS categories, not the federal BKA Gewaltkriminalität definition."
  },
  "features":out_features
}
save(OUT_BERLIN,berlin)
print("BERLIN",matched,"proxy_max",max(f["properties"]["violence_proxy_hz"] for f in out_features))
