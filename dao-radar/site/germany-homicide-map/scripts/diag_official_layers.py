#!/usr/bin/env python3
import io,re,requests
from bs4 import BeautifulSoup
import openpyxl
HEAD={"User-Agent":"Mozilla/5.0 GermanyCrimeMonitor/1.0"}
pages=[
 ("govdata_land","https://www.govdata.de/suche/daten/2025-polizeiliche-kriminalstatistik-t01-grundtabelle-bundeslaender"),
 ("berlin_atlas","https://www.berlin.de/polizei/service/kriminalitaetsatlas/")
]
for label,url in pages:
    try:
        r=requests.get(url,headers=HEAD,timeout=45)
        print("PAGE",label,r.status_code,r.url,len(r.content),r.headers.get("content-type"))
        r.raise_for_status()
    except Exception as e:
        print("PAGE_ERROR",label,repr(e)); continue
    s=BeautifulSoup(r.text,"html.parser")
    links=[]
    for a in s.find_all("a",href=True):
        href=requests.compat.urljoin(r.url,a["href"])
        txt=" ".join(a.get_text(" ",strip=True).split())
        if re.search(r"csv|xlsx|ressource|download|kriminal|bka",txt+" "+href,re.I):
            links.append((txt,href))
    print("LINKS",label,len(links))
    for x in links[:120]: print("L",repr(x[0]),x[1])
    if label=="berlin_atlas":
        for txt,h in links:
            if "xlsx" not in h.lower() and "download" not in txt.lower() and "xlsx" not in txt.lower():
                continue
            try:
                rr=requests.get(h,headers=HEAD,timeout=60)
                print("TRY_XLS",repr(txt),rr.status_code,rr.url,len(rr.content),rr.headers.get("content-type"))
                if rr.ok and rr.content[:2]==b"PK":
                    wb=openpyxl.load_workbook(io.BytesIO(rr.content),read_only=True,data_only=True)
                    print("SHEETS",wb.sheetnames)
                    for ws in wb.worksheets[:12]:
                        print("SHEET",ws.title,ws.max_row,ws.max_column)
                        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,12),values_only=True):
                            print("ROW",ws.title,repr(row[:14]))
                        print("END_SHEET")
                    break
            except Exception as e:
                print("XLS_ERROR",repr(e))
